"""Offline export acceptance tests; create and reopen real Excel/JSON files."""

import json

import pytest
from openpyxl import load_workbook

from GetDataFromAPI.export import changes_between, publish_exports
from GetDataFromAPI.pipeline import run, write_candidate
from GetDataFromAPI.test_pipeline import baseline


def test_excel_json_and_change_files_are_created(tmp_path):
    old = baseline()
    new = [{**old[0], "name": "=external text", "plz": "01067"}]
    assert publish_exports(tmp_path / "pflegestellen.csv", old, new) == 2
    for suffix in ("old", "new", "changes"):
        data = json.loads((tmp_path / f"pflegestellen_{suffix}.json").read_text(encoding="utf-8"))
        workbook = load_workbook(tmp_path / f"pflegestellen_{suffix}.xlsx")
        assert workbook.active.max_row == len(data) + 1
        assert workbook.active.freeze_panes == "A2"
        workbook.close()
    workbook = load_workbook(tmp_path / "pflegestellen_new.xlsx")
    assert workbook.active["C2"].value == "=external text"
    assert workbook.active["C2"].data_type == "s"
    assert workbook.active["E2"].value == "01067"
    workbook.close()


def test_change_detection_ignores_timestamps_and_representation():
    old = baseline()
    new = [{**old[0], "unsere_id": "42", "updated_at": "later", "pflegeberatung": "0"}]
    assert changes_between(old, new) == []
    assert changes_between(old, [])[0]["change"] == "removed"
    assert changes_between([], old)[0]["change"] == "added"


def test_second_import_rotates_last_new_to_old(tmp_path):
    original, output = tmp_path / "baseline.csv", tmp_path / "pflegestellen.csv"
    write_candidate(original, baseline())
    for label in ("First import", "Second import"):
        run(
            "https://source.example",
            original,
            output,
            coordinates=[(1, 1)],
            fetch=lambda *args, label=label: [{"id": "known", "name": label}, {"id": "new", "name": "New"}],
        )
    old = json.loads((tmp_path / "pflegestellen_old.json").read_text(encoding="utf-8"))
    new = json.loads((tmp_path / "pflegestellen_new.json").read_text(encoding="utf-8"))
    changes = json.loads((tmp_path / "pflegestellen_changes.json").read_text(encoding="utf-8"))
    assert old[0]["name"] == "First import" and new[0]["name"] == "Second import"
    assert new[1]["unsere_id"] == old[1]["unsere_id"] == 43
    assert len(changes) == 1 and changes[0]["field"] == "name"


def test_failed_excel_generation_preserves_previous_exports(tmp_path, monkeypatch):
    path = tmp_path / "pflegestellen.csv"
    publish_exports(path, baseline(), baseline())
    before = {p.name: p.read_bytes() for p in tmp_path.iterdir()}

    def fail(*args):
        raise OSError("write error")

    monkeypatch.setattr("GetDataFromAPI.export.write_workbook", fail)
    with pytest.raises(OSError):
        publish_exports(path, baseline(), [{**baseline()[0], "name": "Changed"}])
    assert {p.name: p.read_bytes() for p in tmp_path.iterdir()} == before


def test_publish_failure_rolls_back_every_file(tmp_path, monkeypatch):
    import os

    path = tmp_path / "pflegestellen.csv"
    publish_exports(path, baseline(), baseline())
    before = {p.name: p.read_bytes() for p in tmp_path.iterdir()}
    real_replace = os.replace
    calls = []

    def fail_once(source, target):
        calls.append(1)
        if len(calls) == 4:
            raise OSError("replace failure")
        return real_replace(source, target)

    monkeypatch.setattr("GetDataFromAPI.export.os.replace", fail_once)
    with pytest.raises(OSError):
        publish_exports(path, baseline(), [{**baseline()[0], "name": "Changed"}])
    assert {p.name: p.read_bytes() for p in tmp_path.iterdir()} == before


def test_unknown_nameless_records_are_reported_and_known_records_preserved(tmp_path):
    original, output = tmp_path / "baseline.csv", tmp_path / "pflegestellen.csv"
    records = baseline()
    records[0]["name"] = ""
    write_candidate(original, records)
    run(
        "https://source.example",
        original,
        output,
        coordinates=[(1, 1)],
        fetch=lambda *args: [{"id": "known", "name": None}, {"id": "unknown", "name": None}],
    )
    new = json.loads((tmp_path / "pflegestellen_new.json").read_text(encoding="utf-8"))
    rejected = json.loads((tmp_path / "pflegestellen_rejected.json").read_text(encoding="utf-8"))
    assert len(new) == 1 and new[0]["unsere_id"] == 42
    assert len(rejected) == 1 and rejected[0]["zqp_id"] == "unknown"
