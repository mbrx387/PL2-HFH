"""Generate old/new snapshots and field-level differences from imported records."""

import json
import os
import shutil
import tempfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from GetDataFromAPI.transform import parse_timestamp

CHANGE_COLUMNS = ("change", "unsere_id", "name", "field", "old", "new")


def normalize(records):
    from GetDataFromAPI.pipeline import SERVICE_FIELDS

    result = []
    for record in records:
        row = dict(record)
        row["unsere_id"] = int(row["unsere_id"])
        for field in SERVICE_FIELDS:
            row[field] = int(str(row.get(field, "")).lower() in {"1", "true", "yes", "ja"})
        for field in ("latitude", "longitude"):
            value = row.get(field)
            row[field] = None if value in (None, "") else float(value)
        result.append(row)
    return result


def changes_between(old, new):
    """Ignore bookkeeping timestamps; compare business values by stable app ID."""
    from GetDataFromAPI.pipeline import COLUMNS

    before = {row["unsere_id"]: row for row in normalize(old)}
    after = {row["unsere_id"]: row for row in normalize(new)}
    changes = []
    for identifier in sorted(before.keys() | after.keys()):
        previous, current = before.get(identifier), after.get(identifier)
        name = (current or previous)["name"]
        if previous is None or current is None:
            changes.append(
                dict(
                    zip(
                        CHANGE_COLUMNS,
                        (
                            "added" if previous is None else "removed",
                            identifier,
                            name,
                            "record",
                            json.dumps(previous, ensure_ascii=False) if previous else "",
                            json.dumps(current, ensure_ascii=False) if current else "",
                        ),
                        strict=True,
                    )
                )
            )
            continue
        for field in COLUMNS:
            if field not in {"unsere_id", "updated_at"} and previous.get(field, "") != current.get(field, ""):
                changes.append(
                    dict(
                        zip(
                            CHANGE_COLUMNS,
                            (
                                "modified",
                                identifier,
                                name,
                                field,
                                previous.get(field),
                                current.get(field),
                            ),
                            strict=True,
                        )
                    )
                )
    return changes


def write_workbook(path, rows, columns, title):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = title
    sheet.append(list(columns))
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="24476A")
    for row_index, row in enumerate(rows, start=2):
        sheet.append([row.get(field, "") for field in columns])
        # Source text must remain text, including strings starting with '='.
        for cell in next(sheet.iter_rows(min_row=row_index, max_row=row_index, max_col=len(columns))):
            if columns[cell.column - 1] == "updated_at" and parse_timestamp(cell.value) is not None:
                cell.value = parse_timestamp(cell.value)
                cell.number_format = "dd.mm.yyyy hh:mm:ss"
            elif isinstance(cell.value, str):
                cell.data_type = "s"
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cells in sheet.columns:
        letter = cells[0].column_letter
        field = cells[0].value
        sheet.column_dimensions[letter].width = 18 if field not in {"name", "adresse", "old", "new"} else 45
        if field == "updated_at":
            sheet.column_dimensions[letter].width = 24
        if field in {"plz", "telefon", "zqp_id"}:
            for cell in cells[1:]:
                cell.number_format = "@"
    workbook.save(path)
    workbook.close()


def publish_exports(output_path, old, new, *, rejected=()):
    """Stage all files first, then publish with rollback on ordinary write errors.

    This is not a filesystem transaction across process crashes. A failed
    extraction never reaches this function; a failed export restores old files.
    """
    from GetDataFromAPI.pipeline import COLUMNS, write_candidate

    parent = output_path.parent
    parent.mkdir(parents=True, exist_ok=True)
    stem = output_path.stem
    old, new = normalize(old), normalize(new)
    changes = changes_between(old, new)
    with tempfile.TemporaryDirectory(prefix=".export-", dir=parent) as temporary:
        stage = Path(temporary)
        backup = stage / "backup"
        backup.mkdir()
        write_candidate(stage / output_path.name, new)
        for suffix, rows in (("old", old), ("new", new), ("changes", changes)):
            (stage / f"{stem}_{suffix}.json").write_text(
                json.dumps(rows, ensure_ascii=False, indent=2, allow_nan=False),
                encoding="utf-8",
            )
            write_workbook(
                stage / f"{stem}_{suffix}.xlsx",
                rows,
                CHANGE_COLUMNS if suffix == "changes" else COLUMNS,
                "Changes" if suffix == "changes" else "Pflegestellen",
            )
        (stage / f"{stem}_rejected.json").write_text(
            json.dumps(list(rejected), ensure_ascii=False, indent=2),
            encoding="utf-8",
        )
        write_workbook(stage / f"{stem}_rejected.xlsx", rejected, ("zqp_id", "reason"), "Quality issues")
        files = [path for path in stage.iterdir() if path.is_file()]
        existing = {path.name for path in files if (parent / path.name).exists()}
        for name in existing:
            shutil.copy2(parent / name, backup / name)
        published = []
        try:
            for path in files:
                os.replace(path, parent / path.name)
                published.append(path.name)
        except OSError:
            for name in reversed(published):
                if name in existing:
                    shutil.copy2(backup / name, parent / name)
                else:
                    (parent / name).unlink(missing_ok=True)
            raise
    return len(changes)
